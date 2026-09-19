import io
import pytest
from httpx import AsyncClient, ASGITransport
from app.main import app
from app.models.user import User, UserRole, VerificationStatus
from app.models.document import LegalDocument, DocumentChunk, ProcessingStatus, DocumentType
from app.core.security import get_password_hash


@pytest.mark.asyncio
async def test_valid_pdf_processing(db):
    # Setup Verified Legal Professional
    pro_user = User(
        name="Adv. Kapoor",
        email="kapoor@example.com",
        password_hash=get_password_hash("pass123"),
        role=UserRole.LEGAL_PROFESSIONAL,
        verification_status=VerificationStatus.VERIFIED
    )
    db.add(pro_user)
    db.commit()

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        token = (await ac.post("/api/v1/auth/login", json={"email": "kapoor@example.com", "password": "pass123"})).json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}

        # Create valid PDF bytes
        pdf_bytes = b"%PDF-1.4\n1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R >>\nendobj\n4 0 obj\n<< /Length 55 >>\nstream\nBT /F1 12 Tf 100 700 Td (Supreme Court Judgment Analysis) Tj ET\nendstream\nendobj\nxref\n0 5\n0000000000 65535 f \n0000000009 00000 n \n0000000062 00000 n \n0000000125 00000 n \n0000000225 00000 n \ntrailer\n<< /Size 5 /Root 1 0 R >>\nstartxref\n330\n%%EOF"
        pdf_file = ("court_brief.pdf", io.BytesIO(pdf_bytes), "application/pdf")

        response = await ac.post(
            "/api/v1/documents/upload",
            files={"file": pdf_file},
            data={"document_type": "case_brief", "jurisdiction": "Supreme Court of India", "matter_id": "MAT-2026-001"},
            headers=headers
        )

    assert response.status_code == 201
    data = response.json()
    assert data["filename"] == "court_brief.pdf"
    assert data["file_type"] == "pdf"
    assert data["processing_status"] == ProcessingStatus.COMPLETED.value
    assert data["chunk_count"] >= 1
    assert "file_hash" in data


@pytest.mark.asyncio
async def test_valid_docx_processing(db):
    pro_user = User(
        name="Adv. Roy",
        email="roy@example.com",
        password_hash=get_password_hash("pass123"),
        role=UserRole.LEGAL_PROFESSIONAL,
        verification_status=VerificationStatus.VERIFIED
    )
    db.add(pro_user)
    db.commit()

    # Create dummy docx file
    import docx
    doc_io = io.BytesIO()
    doc = docx.Document()
    doc.add_heading("Commercial Lease Agreement", level=1)
    doc.add_paragraph("This agreement is entered into between Lessor and Lessee regarding commercial premises.")
    doc.save(doc_io)
    doc_io.seek(0)

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        token = (await ac.post("/api/v1/auth/login", json={"email": "roy@example.com", "password": "pass123"})).json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}

        docx_file = ("contract.docx", doc_io, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        response = await ac.post(
            "/api/v1/documents/upload",
            files={"file": docx_file},
            data={"document_type": "contract"},
            headers=headers
        )

    assert response.status_code == 201
    data = response.json()
    assert data["filename"] == "contract.docx"
    assert data["file_type"] == "docx"
    assert data["processing_status"] == ProcessingStatus.COMPLETED.value
    assert data["chunk_count"] >= 1


@pytest.mark.asyncio
async def test_valid_txt_processing(db):
    pro_user = User(
        name="Adv. Mehta",
        email="mehta@example.com",
        password_hash=get_password_hash("pass123"),
        role=UserRole.LEGAL_PROFESSIONAL,
        verification_status=VerificationStatus.VERIFIED
    )
    db.add(pro_user)
    db.commit()

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        token = (await ac.post("/api/v1/auth/login", json={"email": "mehta@example.com", "password": "pass123"})).json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}

        txt_bytes = b"Statutory Provisions: Section 420 IPC Penalties and Compliance Guidelines."
        txt_file = ("statute.txt", io.BytesIO(txt_bytes), "text/plain")

        response = await ac.post(
            "/api/v1/documents/upload",
            files={"file": txt_file},
            data={"document_type": "statute"},
            headers=headers
        )

    assert response.status_code == 201
    data = response.json()
    assert data["filename"] == "statute.txt"
    assert data["file_type"] == "txt"
    assert data["processing_status"] == ProcessingStatus.COMPLETED.value
    assert data["chunk_count"] >= 1


@pytest.mark.asyncio
async def test_invalid_extension(db):
    pro_user = User(
        name="Adv. Pro",
        email="pro_ext@example.com",
        password_hash=get_password_hash("pass123"),
        role=UserRole.LEGAL_PROFESSIONAL,
        verification_status=VerificationStatus.VERIFIED
    )
    db.add(pro_user)
    db.commit()

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        token = (await ac.post("/api/v1/auth/login", json={"email": "pro_ext@example.com", "password": "pass123"})).json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}

        # Try uploading unsupported file type (.png or .exe)
        bad_file = ("image.png", io.BytesIO(b"png_image_bytes"), "image/png")
        response = await ac.post("/api/v1/documents/upload", files={"file": bad_file}, headers=headers)

    assert response.status_code == 400
    assert "Unsupported document format" in response.json()["detail"]


@pytest.mark.asyncio
async def test_oversized_file(db):
    pro_user = User(
        name="Adv. Pro",
        email="pro_size@example.com",
        password_hash=get_password_hash("pass123"),
        role=UserRole.LEGAL_PROFESSIONAL,
        verification_status=VerificationStatus.VERIFIED
    )
    db.add(pro_user)
    db.commit()

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        token = (await ac.post("/api/v1/auth/login", json={"email": "pro_size@example.com", "password": "pass123"})).json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}

        # Oversized file (> 15MB)
        large_bytes = b"0" * (16 * 1024 * 1024)
        large_file = ("large.txt", io.BytesIO(large_bytes), "text/plain")
        response = await ac.post("/api/v1/documents/upload", files={"file": large_file}, headers=headers)

    assert response.status_code == 400
    assert "exceeds maximum allowed limit" in response.json()["detail"]


@pytest.mark.asyncio
async def test_unauthorized_document_access(db):
    public_user = User(
        name="Public User",
        email="public_doc@example.com",
        password_hash=get_password_hash("pass123"),
        role=UserRole.PUBLIC_USER
    )
    db.add(public_user)
    db.commit()

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        token = (await ac.post("/api/v1/auth/login", json={"email": "public_doc@example.com", "password": "pass123"})).json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}

        txt_file = ("doc.txt", io.BytesIO(b"Sample Text"), "text/plain")
        response = await ac.post("/api/v1/documents/upload", files={"file": txt_file}, headers=headers)

    assert response.status_code == 403
    assert "Permission denied" in response.json()["detail"]


@pytest.mark.asyncio
async def test_ownership_checks(db):
    user_a = User(name="User A", email="usera@example.com", password_hash=get_password_hash("pass123"), role=UserRole.LEGAL_PROFESSIONAL, verification_status=VerificationStatus.VERIFIED)
    user_b = User(name="User B", email="userb@example.com", password_hash=get_password_hash("pass123"), role=UserRole.LEGAL_PROFESSIONAL, verification_status=VerificationStatus.VERIFIED)
    admin = User(name="Admin", email="admin_doc@example.com", password_hash=get_password_hash("pass123"), role=UserRole.ADMIN)
    db.add_all([user_a, user_b, admin])
    db.commit()

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        token_a = (await ac.post("/api/v1/auth/login", json={"email": "usera@example.com", "password": "pass123"})).json()["access_token"]
        token_b = (await ac.post("/api/v1/auth/login", json={"email": "userb@example.com", "password": "pass123"})).json()["access_token"]
        token_admin = (await ac.post("/api/v1/auth/login", json={"email": "admin_doc@example.com", "password": "pass123"})).json()["access_token"]

        # User A uploads document
        upload_resp = await ac.post(
            "/api/v1/documents/upload",
            files={"file": ("doc_a.txt", io.BytesIO(b"Private Legal Brief A"), "text/plain")},
            headers={"Authorization": f"Bearer {token_a}"}
        )
        doc_id = upload_resp.json()["id"]

        # User B attempts to view User A's document details -> 403 Forbidden
        b_resp = await ac.get(f"/api/v1/documents/{doc_id}", headers={"Authorization": f"Bearer {token_b}"})
        assert b_resp.status_code == 403

        # User A views own document details -> 200 OK
        a_resp = await ac.get(f"/api/v1/documents/{doc_id}", headers={"Authorization": f"Bearer {token_a}"})
        assert a_resp.status_code == 200

        # Admin views User A's document details -> 200 OK
        admin_resp = await ac.get(f"/api/v1/documents/{doc_id}", headers={"Authorization": f"Bearer {token_admin}"})
        assert admin_resp.status_code == 200


@pytest.mark.asyncio
async def test_extraction_failure_handling(db):
    pro_user = User(name="Adv. Fail", email="fail@example.com", password_hash=get_password_hash("pass123"), role=UserRole.LEGAL_PROFESSIONAL, verification_status=VerificationStatus.VERIFIED)
    db.add(pro_user)
    db.commit()

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        token = (await ac.post("/api/v1/auth/login", json={"email": "fail@example.com", "password": "pass123"})).json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}

        # Corrupted PDF content
        corrupt_file = ("corrupt.pdf", io.BytesIO(b"NOT_A_VALID_PDF_STRUCTURE_INVALID_HEADER"), "application/pdf")
        response = await ac.post("/api/v1/documents/upload", files={"file": corrupt_file}, headers=headers)

    assert response.status_code == 201
    data = response.json()
    assert data["processing_status"] == ProcessingStatus.FAILED.value
    assert "Processing Failure" in data["error_message"]
